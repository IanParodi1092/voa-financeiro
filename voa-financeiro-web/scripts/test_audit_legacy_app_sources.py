from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))

from audit_legacy_app_sources import (  # noqa: E402
    audit,
    collect,
    find_app_source,
    is_base_currency,
    is_system_column,
    normalize,
    read_app_sources,
    read_dataverse_entities,
    read_field_references,
    read_workbook_headers,
    render_report,
    resolve_sheet,
    resolve_variant_sheets,
    sheet_coverage,
    suggestions,
)

ALUNO_MAPPING = {
    "cr04f_tbl_alunoid": "TBL_ALUNO",
    "cr04f_titulo": "Title",
    "cr04f_nomedoaluno": "Aluno.Nome",
    "cr04f_escoladoaluno": "Aluno_Escola",
    "cr04f_valor_base": "Aluno_Valor (Base)",
    "cr04f_datadecriacao": "Created",
    "cr04f_possuianexos": "Attachments",
}
FANTASMA_MAPPING = {"cr04f_tbl_fantasmaid": "TBL_FANTASMA", "cr04f_nome": "Fantasma_Nome"}
ENTITY_XML = """<?xml version="1.0" encoding="utf-8"?>
<Entity>
  <Name LocalizedName="TBL_ALUNO" OriginalName="TBL_ALUNO">cr04f_tbl_aluno</Name>
  <attribute PhysicalName="cr04f_titulo">
    <displaynames><displayname description="Title" languagecode="1033" /></displaynames>
  </attribute>
  <attribute PhysicalName="cr04f_nomedoaluno">
    <displaynames><displayname description="Aluno.Nome" languagecode="1033" /></displaynames>
  </attribute>
  <attribute PhysicalName="cr04f_escoladoaluno">
    <displaynames><displayname description="Aluno_Escola" languagecode="1033" /></displaynames>
  </attribute>
  <attribute PhysicalName="cr04f_valor_base">
    <displaynames><displayname description="Aluno_Valor (Base)" languagecode="1033" /></displaynames>
  </attribute>
  <attribute PhysicalName="CreatedOn">
    <displaynames><displayname description="Data de Criacao" languagecode="1033" /></displaynames>
  </attribute>
</Entity>
"""


def write_fixture(root: Path) -> Path:
    source = root / "app"
    (source / "DataSources").mkdir(parents=True)
    (source / "Src").mkdir(parents=True)
    (source / "DataSources" / "TBL_ALUNO.json").write_text(
        json.dumps(
            [
                {
                    "Type": "NativeCDSDataSourceInfo",
                    "Name": "TBL_ALUNO",
                    "LogicalName": "cr04f_tbl_aluno",
                    "NativeCDSDataSourceInfoNameMapping": ALUNO_MAPPING,
                }
            ]
        ),
        encoding="utf-8",
    )
    (source / "DataSources" / "TBL_FANTASMA.json").write_text(
        json.dumps(
            [
                {
                    "Type": "NativeCDSDataSourceInfo",
                    "Name": "TBL_FANTASMA",
                    "LogicalName": "cr04f_tbl_fantasma",
                    "NativeCDSDataSourceInfoNameMapping": FANTASMA_MAPPING,
                }
            ]
        ),
        encoding="utf-8",
    )
    (source / "DataSources" / "ImportarExtrato.json").write_text(
        json.dumps(
            [
                {
                    "Type": "ConnectedDataSourceInfo",
                    "Name": "ImportarExtrato",
                    "DatasetName": "me,b!planilha",
                    "ConnectedDataSourceInfoNameMapping": {"Valor": "Valor", "Conta": "Conta"},
                }
            ]
        ),
        encoding="utf-8",
    )
    (source / "DataSources" / "CriarBoleto.json").write_text(
        json.dumps([{"Type": "ServiceInfo", "Name": "CriarBoleto", "ServiceKind": "ConnectedWadl"}]),
        encoding="utf-8",
    )
    (source / "Src" / "Screen_Teste.fx.yaml").write_text(
        "\n".join(
            [
                "Screen_Teste As screen:",
                "    Items: =Filter(TBL_FANTASMA, Fantasma_Nome = \"Escola\")",
                "    Label1: ='Aluno.Nome' & ThisItem.'Aluno_Escola'",
                "    Label2: =ThisItem.'Aluno_Valor (Base)'",
            ]
        ),
        encoding="utf-8",
    )
    return source


def write_workbook(root: Path) -> Path:
    path = root / "VOA.xlsx"
    book = openpyxl.Workbook()
    aluno = book.active
    aluno.title = "TBL_ALUNO"
    aluno.append(["Title", "Aluno.Nome"])
    aluno.append(["Titulo", "Maria"])
    sobra = book.create_sheet("TBL_SOBRA")
    sobra.append(["ColunaA", "ColunaB"])
    book.save(path)
    return path


def write_solution(root: Path) -> Path:
    entities = root / "Entities" / "cr04f_TBL_ALUNO"
    entities.mkdir(parents=True)
    (entities / "Entity.xml").write_text(ENTITY_XML, encoding="utf-8")
    return root / "Entities"


def build_report(root: Path) -> dict:
    source = write_fixture(root)
    workbook = write_workbook(root)
    solution = write_solution(root)
    return collect(source, workbook, solution)


class HelperTests(unittest.TestCase):
    def test_normalize_folds_punctuation(self) -> None:
        self.assertEqual(normalize("TBL_PAGAMENTO (2)"), "tblpagamento2")
        self.assertEqual(normalize("Aluno.Nome"), "alunonome")

    def test_system_and_base_currency_buckets(self) -> None:
        self.assertTrue(is_system_column("Created By"))
        self.assertTrue(is_system_column("OData__ColorTag"))
        self.assertTrue(is_system_column("OData__UIVersionString"))
        self.assertTrue(is_system_column("ID.1"))
        self.assertFalse(is_system_column("Aluno.Nome"))
        self.assertFalse(is_system_column("FI"))
        self.assertTrue(is_base_currency("Pagamento_ValorParcelaTotal (Base)"))
        self.assertTrue(is_base_currency("ExchangeRate"))
        self.assertFalse(is_base_currency("CONTA_NOME"))

    def test_resolve_sheet_exact_and_loose(self) -> None:
        headers = {"TBL_ALUNO": ["Title"], "TBL_SOBRA": []}
        self.assertEqual(resolve_sheet("TBL_ALUNO", headers), ("TBL_ALUNO", True))
        self.assertEqual(resolve_sheet("cr04f_TBL_ALUNO", headers), ("TBL_ALUNO", False))

    def test_resolve_variant_sheets_requires_a_suffix_and_a_header_row(self) -> None:
        headers = {
            "TBL_EXTRATO": [],
            "TBL_EXTRATO_FINANCEIRO": ["EXTRATO_VALOR"],
            "TBL_EXTRATO_VAZIA": [],
            "TBL_CLIENTEENDTEL": ["CLIENTEENDTEL_VALOR"],
        }
        taken = {"tblextrato", "tblcliente", "tblclienteendtel"}
        self.assertEqual(
            resolve_variant_sheets("TBL_EXTRATO", headers, taken),
            ["TBL_EXTRATO_FINANCEIRO"],
        )
        # `TBL_CLIENTEENDTEL` e a aba canonica de outra tabela, nao uma variante de `TBL_CLIENTE`.
        self.assertEqual(resolve_variant_sheets("TBL_CLIENTE", headers, taken), [])

    def test_sheet_coverage_counts_only_folded_matches(self) -> None:
        headers = {"TBL_EXTRATO_FINANCEIRO": ["EXTRATO_VALOR", "extrato_data_venc"]}
        self.assertEqual(sheet_coverage("TBL_EXTRATO_FINANCEIRO", headers, ["EXTRATO_VALOR", "EXTRATO_EMPRESA_ID"]), 1)
        self.assertEqual(sheet_coverage(None, headers, ["EXTRATO_VALOR"]), 0)

    def test_suggestions_point_to_the_typo_variant(self) -> None:
        self.assertEqual(suggestions("TBL_EMAILLISTACONTATO", ["TBL_EMAILISTACONTATO"]), ["TBL_EMAILISTACONTATO"])

    def test_find_app_source_uses_the_minimum_threshold(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = write_fixture(root)
            self.assertEqual(find_app_source(root, minimum_sources=2), source)
            with self.assertRaises(FileNotFoundError):
                find_app_source(root)


class ReaderTests(unittest.TestCase):
    def test_read_app_sources_skips_the_primary_key_entry(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = write_fixture(Path(tmp))
            sources = {entry["name"]: entry for entry in read_app_sources(source)}
            self.assertEqual(set(sources), {"TBL_ALUNO", "TBL_FANTASMA", "ImportarExtrato", "CriarBoleto"})
            self.assertNotIn("TBL_ALUNO", sources["TBL_ALUNO"]["fields"].values())
            self.assertEqual(sources["TBL_ALUNO"]["fields"]["cr04f_nomedoaluno"], "Aluno.Nome")
            self.assertEqual(sources["ImportarExtrato"]["type"], "ConnectedDataSourceInfo")

    def test_read_field_references_counts_both_reference_styles(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = write_fixture(Path(tmp))
            universe = {"Fantasma_Nome": {"TBL_FANTASMA"}, "Aluno_Escola": {"TBL_ALUNO"}, "Aluno.Nome": {"TBL_ALUNO"}}
            counts, where = read_field_references(source, universe)
            self.assertEqual(counts["TBL_FANTASMA|Fantasma_Nome"], 1)
            self.assertEqual(counts["TBL_ALUNO|Aluno_Escola"], 1)
            self.assertEqual(counts["TBL_ALUNO|Aluno.Nome"], 1)
            self.assertEqual(where["TBL_ALUNO|Aluno.Nome"], {"Screen_Teste.fx.yaml"})

    def test_read_dataverse_entities_uses_display_names(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            entities = read_dataverse_entities(write_solution(Path(tmp)))
            self.assertEqual(list(entities), ["TBL_ALUNO"])
            labels = {attribute["label"] for attribute in entities["TBL_ALUNO"]["attributes"]}
            self.assertIn("Aluno_Escola", labels)
            self.assertIn("Data de Criacao", labels)

    def test_read_workbook_headers_reads_row_one(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            headers = read_workbook_headers(write_workbook(Path(tmp)))
            self.assertEqual(headers["TBL_ALUNO"], ["Title", "Aluno.Nome"])
            self.assertEqual(headers["TBL_SOBRA"], ["ColunaA", "ColunaB"])


class AuditTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)
        self.source = write_fixture(self.root)
        self.workbook = write_workbook(self.root)
        self.solution = write_solution(self.root)
        self.report = collect(self.source, self.workbook, self.solution)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_counts_every_input(self) -> None:
        counts = self.report["counts"]
        self.assertEqual(counts["app_data_sources"], 4)
        self.assertEqual(counts["app_tables"], 2)
        self.assertEqual(counts["app_external_sources"], 2)
        self.assertEqual(counts["workbook_sheets"], 2)
        self.assertEqual(counts["dataverse_entities"], 1)

    def test_table_without_sheet_lists_the_used_business_fields(self) -> None:
        self.assertEqual([item["table"] for item in self.report["tables_without_sheet"]], ["TBL_FANTASMA"])
        item = self.report["tables_without_sheet"][0]
        self.assertEqual(item["used"], [("Fantasma_Nome", 1)])
        self.assertEqual(item["used_total"], 1)

    def test_field_gaps_separate_business_system_and_base_currency(self) -> None:
        gap = next(item for item in self.report["field_gaps"] if item["table"] == "TBL_ALUNO")
        self.assertEqual([field["field"] for field in gap["business"]], ["Aluno_Escola"])
        self.assertEqual([field["field"] for field in gap["base_currency"]], ["Aluno_Valor (Base)"])
        self.assertEqual(gap["business"][0]["files"], ["Screen_Teste.fx.yaml"])
        rendered = {field["field"] for field in gap["business"] + gap["base_currency"]}
        self.assertNotIn("Created", rendered)
        self.assertNotIn("Attachments", rendered)
        self.assertNotIn("Title", rendered)

    def test_extra_sheet_has_no_entity(self) -> None:
        self.assertEqual([item["sheet"] for item in self.report["extra_sheets"]], ["TBL_SOBRA"])
        self.assertIsNone(self.report["extra_sheets"][0]["entity"])

    def test_external_sources_keep_their_column_list(self) -> None:
        external = {item["name"]: item for item in self.report["external_sources"]}
        self.assertEqual(external["ImportarExtrato"]["type"], "ConnectedDataSourceInfo")
        self.assertEqual(external["ImportarExtrato"]["fields"], ["Conta", "Valor"])
        self.assertEqual(external["CriarBoleto"]["detail"], "ConnectedWadl")

    def test_variant_sheet_covers_the_table_whose_canonical_sheet_is_empty(self) -> None:
        path = self.root / "VARIANTE.xlsx"
        book = openpyxl.Workbook()
        empty = book.active
        empty.title = "TBL_ALUNO"
        financeiro = book.create_sheet("TBL_ALUNO_FINANCEIRO")
        financeiro.append(["Title", "Aluno.Nome"])
        fantasma = book.create_sheet("TBL_FANTASMA_FINANCEIRO")
        fantasma.append(["Fantasma_Nome"])
        book.save(path)

        report = collect(self.source, path, self.solution)

        self.assertEqual(report["tables_without_sheet"], [])
        self.assertEqual(report["counts"]["variant_sheets"], 2)
        variants = {item["table"]: item for item in report["variant_sheets"]}
        self.assertEqual(variants["TBL_ALUNO"]["sheet"], "TBL_ALUNO")
        self.assertEqual(variants["TBL_ALUNO"]["variant"], "TBL_ALUNO_FINANCEIRO")
        self.assertEqual(variants["TBL_ALUNO"]["covered"], 1)
        self.assertEqual(variants["TBL_ALUNO"]["used_total"], 3)
        self.assertIsNone(variants["TBL_FANTASMA"]["sheet"])
        gap = next(item for item in report["field_gaps"] if item["table"] == "TBL_ALUNO")
        self.assertEqual([field["field"] for field in gap["business"]], ["Aluno_Escola"])
        self.assertEqual([field["field"] for field in gap["base_currency"]], ["Aluno_Valor (Base)"])
        self.assertTrue(gap["variant"])
        self.assertEqual(report["extra_sheets"], [])

        text = render_report(report, {"app_source": "app", "workbook": "book", "solution": "sol"})
        self.assertIn("## 3b. Abas irmas que entregam o conteudo da tabela", text)
        self.assertIn("| TBL_ALUNO | TBL_ALUNO | TBL_ALUNO_FINANCEIRO | 1/3 |", text)

    def test_loose_sheet_name_is_reported_as_divergence(self) -> None:
        path = self.root / "OUTRO.xlsx"
        book = openpyxl.Workbook()
        book.active.title = "cr04f_TBL_ALUNO"
        book.active.append(["Title"])
        book.save(path)
        report = collect(self.source, path, self.solution)
        self.assertEqual(report["name_mismatch"], [{"table": "TBL_ALUNO", "sheet": "cr04f_TBL_ALUNO"}])

    def test_render_report_marks_every_section(self) -> None:
        text = render_report(
            self.report,
            {"app_source": "app", "workbook": "book", "solution": "sol"},
        )
        self.assertIn("## 1. Tabelas que a app le", text)
        self.assertIn("| TBL_FANTASMA | `cr04f_tbl_fantasma` |", text)
        self.assertIn("| TBL_ALUNO | TBL_ALUNO | `Aluno_Escola` | 1 | Screen_Teste.fx.yaml |", text)
        self.assertIn("| TBL_SOBRA | 2 | - |", text)
        self.assertIn("## 3b. Abas irmas que entregam o conteudo da tabela", text)
        self.assertIn("ImportarExtrato", text)
        self.assertIn("somente leitura", text)


if __name__ == "__main__":
    unittest.main()


